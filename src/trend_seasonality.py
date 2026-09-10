"""
Trend and seasonality analysis for time-series forecasting prep.

Method (kept deliberately simple to interpret):
  1. Resample raw 10-min readings to hourly means (base series for everything below).
  2. Fit ONE linear trend per zone (ordinary least squares on hours-since-start).
     This single trend line is expressed at 3 zoom levels for the trend line charts:
     yearly (daily-mean view), weekly (weekly-mean view), daily (raw hourly view).
  3. Detrend the hourly series (value - fitted trend) and average the residual by
     month / day-of-week / hour-of-day to get the seasonality bar charts.

Outputs:
  reports/figures/trend/trend_yearly.png, trend_weekly.png, trend_daily_hourly.png
  reports/figures/seasonality/seasonality_yearly_monthly.png,
                              seasonality_weekly_dow.png,
                              seasonality_daily_hourly.png
  reports/03_trend_seasonality_report.xlsx  (Trend sheet, Seasonality sheet, ReadMe sheet)

Run: python src/trend_seasonality.py
"""

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

DATA_PATH = Path("data/processed/cleaned_energy_data.csv")
TREND_DIR = Path("reports/figures/trend")
SEASON_DIR = Path("reports/figures/seasonality")
XLSX_PATH = Path("reports/03_trend_seasonality_report.xlsx")

ZONES = {
    "PowerConsumption_Zone1": "Zone 1",
    "PowerConsumption_Zone2": "Zone 2",
    "PowerConsumption_Zone3": "Zone 3",
}
DOW_LABELS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
MONTH_LABELS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

TREND_DIR.mkdir(parents=True, exist_ok=True)
SEASON_DIR.mkdir(parents=True, exist_ok=True)


def load_hourly():
    df = pd.read_csv(DATA_PATH, parse_dates=["Datetime"]).set_index("Datetime")
    return df[list(ZONES.keys())].resample("h").mean()


def fit_trend(series):
    x = np.arange(len(series))
    slope, intercept = np.polyfit(x, series.values, 1)
    fitted = slope * x + intercept
    return slope, intercept, pd.Series(fitted, index=series.index)


def plot_trend_charts(hdf, trends):
    # --- Yearly view: daily means + trend ---
    fig, axes = plt.subplots(3, 1, figsize=(13, 9), sharex=True)
    for ax, (col, label) in zip(axes, ZONES.items()):
        daily = hdf[col].resample("D").mean()
        trend_daily = trends[col]["fitted"].resample("D").mean()
        ax.plot(daily.index, daily.values, color="steelblue", linewidth=1, label="Daily mean")
        ax.plot(trend_daily.index, trend_daily.values, color="crimson", linestyle="--", linewidth=2, label="Trend")
        slope_day = trends[col]["slope"] * 24
        ax.set_title(f"{label}: Yearly trend (slope = {slope_day:+.1f} kW/day)")
        ax.legend(fontsize=8)
        ax.set_ylabel("kW")
    plt.xlabel("Date")
    plt.tight_layout()
    plt.savefig(TREND_DIR / "trend_yearly.png", dpi=110)
    plt.close()

    # --- Weekly view: weekly means + trend ---
    fig, axes = plt.subplots(3, 1, figsize=(13, 9), sharex=True)
    for ax, (col, label) in zip(axes, ZONES.items()):
        weekly = hdf[col].resample("W").mean()
        trend_weekly = trends[col]["fitted"].resample("W").mean()
        ax.plot(range(len(weekly)), weekly.values, color="seagreen", linewidth=1.3, marker="o", markersize=3, label="Weekly mean")
        ax.plot(range(len(trend_weekly)), trend_weekly.values, color="crimson", linestyle="--", linewidth=2, label="Trend")
        slope_week = trends[col]["slope"] * 24 * 7
        ax.set_title(f"{label}: Weekly trend (slope = {slope_week:+.1f} kW/week)")
        ax.legend(fontsize=8)
        ax.set_ylabel("kW")
    plt.xlabel("Week number")
    plt.tight_layout()
    plt.savefig(TREND_DIR / "trend_weekly.png", dpi=110)
    plt.close()

    # --- Daily/hourly view: raw hourly series + trend ---
    fig, axes = plt.subplots(3, 1, figsize=(13, 9), sharex=True)
    for ax, (col, label) in zip(axes, ZONES.items()):
        ax.plot(hdf.index, hdf[col].values, color="gray", linewidth=0.3, label="Hourly mean")
        ax.plot(trends[col]["fitted"].index, trends[col]["fitted"].values, color="crimson", linestyle="--", linewidth=2, label="Trend")
        slope_hr = trends[col]["slope"]
        ax.set_title(f"{label}: Hourly-resolution trend (slope = {slope_hr:+.3f} kW/hour)")
        ax.legend(fontsize=8)
        ax.set_ylabel("kW")
    plt.xlabel("Date")
    plt.tight_layout()
    plt.savefig(TREND_DIR / "trend_daily_hourly.png", dpi=110)
    plt.close()


def plot_seasonality_charts(hdf, trends):
    detrended = pd.DataFrame({col: hdf[col] - trends[col]["fitted"] for col in ZONES})

    # --- Yearly seasonality: monthly bars ---
    fig, axes = plt.subplots(3, 1, figsize=(11, 9), sharex=True)
    monthly_summary = {}
    for ax, (col, label) in zip(axes, ZONES.items()):
        by_month = detrended[col].groupby(detrended.index.month).mean()
        by_month.index = [MONTH_LABELS[m - 1] for m in by_month.index]
        colors = ["crimson" if v > 0 else "steelblue" for v in by_month.values]
        ax.bar(by_month.index, by_month.values, color=colors)
        ax.axhline(0, color="black", linewidth=0.8)
        ax.set_title(f"{label}: Yearly seasonality (detrended, by month)")
        ax.set_ylabel("kW deviation")
        monthly_summary[label] = by_month
    plt.tight_layout()
    plt.savefig(SEASON_DIR / "seasonality_yearly_monthly.png", dpi=110)
    plt.close()

    # --- Weekly seasonality: day-of-week bars ---
    fig, axes = plt.subplots(3, 1, figsize=(9, 9), sharex=True)
    dow_summary = {}
    for ax, (col, label) in zip(axes, ZONES.items()):
        by_dow = detrended[col].groupby(detrended.index.dayofweek).mean()
        by_dow.index = DOW_LABELS
        colors = ["crimson" if v > 0 else "steelblue" for v in by_dow.values]
        ax.bar(by_dow.index, by_dow.values, color=colors)
        ax.axhline(0, color="black", linewidth=0.8)
        ax.set_title(f"{label}: Weekly seasonality (detrended, by day-of-week)")
        ax.set_ylabel("kW deviation")
        dow_summary[label] = by_dow
    plt.tight_layout()
    plt.savefig(SEASON_DIR / "seasonality_weekly_dow.png", dpi=110)
    plt.close()

    # --- Daily seasonality: hour-of-day bars (24h) ---
    fig, axes = plt.subplots(3, 1, figsize=(11, 9), sharex=True)
    hour_summary = {}
    for ax, (col, label) in zip(axes, ZONES.items()):
        by_hour = detrended[col].groupby(detrended.index.hour).mean()
        colors = ["crimson" if v > 0 else "steelblue" for v in by_hour.values]
        ax.bar(by_hour.index, by_hour.values, color=colors)
        ax.axhline(0, color="black", linewidth=0.8)
        ax.set_title(f"{label}: Daily seasonality (detrended, 24h profile)")
        ax.set_ylabel("kW deviation")
        ax.set_xticks(range(0, 24, 2))
        hour_summary[label] = by_hour
    plt.xlabel("Hour of day")
    plt.tight_layout()
    plt.savefig(SEASON_DIR / "seasonality_daily_hourly.png", dpi=110)
    plt.close()

    return monthly_summary, dow_summary, hour_summary


def build_excel(trends, monthly_summary, dow_summary, hour_summary, hdf):
    wb = Workbook()

    # --- ReadMe sheet ---
    ws = wb.active
    ws.title = "ReadMe"
    ws["A1"] = "Trend & Seasonality Report: Method"
    ws["A1"].font = Font(bold=True, size=14)
    lines = [
        "",
        "Base series: raw 10-minute readings resampled to hourly means.",
        "Trend: one linear regression fit per zone (value ~ hours since start). Kept simple so the",
        "slope is a single interpretable number, shown at 3 zoom levels on the Trend sheet:",
        "  - Yearly view  = daily means with the trend line, slope reported in kW/day",
        "  - Weekly view  = weekly means with the trend line, slope reported in kW/week",
        "  - Daily view   = raw hourly means with the trend line, slope reported in kW/hour",
        "(All three are the same underlying trend line, just viewed at different time resolutions.)",
        "",
        "Seasonality: detrended residual (hourly value minus fitted trend) averaged by:",
        "  - Month (Jan-Dec)      -> yearly seasonality",
        "  - Day of week (Mon-Sun) -> weekly seasonality",
        "  - Hour of day (0-23)    -> daily seasonality (24h profile)",
        "Positive bars = above the long-term trend for that period; negative = below trend.",
        "",
        "Files: reports/figures/trend/*.png and reports/figures/seasonality/*.png",
    ]
    for i, line in enumerate(lines, start=2):
        ws[f"A{i}"] = line
    ws.column_dimensions["A"].width = 100

    # --- Trend sheet ---
    ws_t = wb.create_sheet("Trend")
    ws_t["A1"] = "Trend Summary"
    ws_t["A1"].font = Font(bold=True, size=14)
    headers = ["Zone", "Slope (kW/hour)", "Slope (kW/day)", "Slope (kW/week)", "Net change over year (kW)", "Net change (%)"]
    for j, h in enumerate(headers, start=1):
        c = ws_t.cell(row=3, column=j, value=h)
        c.font = Font(bold=True)
    row = 4
    for col, label in ZONES.items():
        slope = trends[col]["slope"]
        mean_val = hdf[col].mean()
        n_hours = len(hdf)
        net_change = slope * n_hours
        pct_change = 100 * net_change / mean_val
        ws_t.cell(row=row, column=1, value=label)
        ws_t.cell(row=row, column=2, value=round(slope, 4))
        ws_t.cell(row=row, column=3, value=round(slope * 24, 2))
        ws_t.cell(row=row, column=4, value=round(slope * 24 * 7, 2))
        ws_t.cell(row=row, column=5, value=round(net_change, 0))
        ws_t.cell(row=row, column=6, value=round(pct_change, 1))
        row += 1
    for j in range(1, 7):
        ws_t.column_dimensions[get_column_letter(j)].width = 22

    img_row = row + 2
    for fname, title in [
        ("trend_yearly.png", "Yearly trend (daily means)"),
        ("trend_weekly.png", "Weekly trend (weekly means)"),
        ("trend_daily_hourly.png", "Hourly-resolution trend"),
    ]:
        ws_t.cell(row=img_row, column=1, value=title).font = Font(bold=True, italic=True)
        img = XLImage(str(TREND_DIR / fname))
        img.width, img.height = 780, 540
        ws_t.add_image(img, f"A{img_row + 1}")
        img_row += 30

    # --- Seasonality sheet ---
    ws_s = wb.create_sheet("Seasonality")
    ws_s["A1"] = "Seasonality Summary (detrended)"
    ws_s["A1"].font = Font(bold=True, size=14)
    headers = ["Zone", "Granularity", "Peak bin", "Peak value (kW dev)", "Trough bin", "Trough value (kW dev)", "Amplitude (kW)"]
    for j, h in enumerate(headers, start=1):
        c = ws_s.cell(row=3, column=j, value=h)
        c.font = Font(bold=True)
    row = 4
    for label in ZONES.values():
        for gran_name, summary in [("Yearly (month)", monthly_summary), ("Weekly (day-of-week)", dow_summary), ("Daily (hour-of-day)", hour_summary)]:
            s = summary[label]
            peak_bin, peak_val = s.idxmax(), s.max()
            trough_bin, trough_val = s.idxmin(), s.min()
            ws_s.cell(row=row, column=1, value=label)
            ws_s.cell(row=row, column=2, value=gran_name)
            ws_s.cell(row=row, column=3, value=str(peak_bin))
            ws_s.cell(row=row, column=4, value=round(peak_val, 1))
            ws_s.cell(row=row, column=5, value=str(trough_bin))
            ws_s.cell(row=row, column=6, value=round(trough_val, 1))
            ws_s.cell(row=row, column=7, value=round(peak_val - trough_val, 1))
            row += 1
    for j in range(1, 8):
        ws_s.column_dimensions[get_column_letter(j)].width = 20

    img_row = row + 2
    for fname, title in [
        ("seasonality_yearly_monthly.png", "Yearly seasonality (by month, detrended)"),
        ("seasonality_weekly_dow.png", "Weekly seasonality (by day-of-week, detrended)"),
        ("seasonality_daily_hourly.png", "Daily seasonality (24h profile, detrended)"),
    ]:
        ws_s.cell(row=img_row, column=1, value=title).font = Font(bold=True, italic=True)
        img = XLImage(str(SEASON_DIR / fname))
        img.width, img.height = 700, 540
        ws_s.add_image(img, f"A{img_row + 1}")
        img_row += 30

    XLSX_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX_PATH)


if __name__ == "__main__":
    hdf = load_hourly()

    trends = {}
    for col in ZONES:
        slope, intercept, fitted = fit_trend(hdf[col])
        trends[col] = {"slope": slope, "intercept": intercept, "fitted": fitted}

    plot_trend_charts(hdf, trends)
    monthly_summary, dow_summary, hour_summary = plot_seasonality_charts(hdf, trends)
    build_excel(trends, monthly_summary, dow_summary, hour_summary, hdf)

    print("Trend & seasonality analysis complete.")
    print(f"Trend charts -> {TREND_DIR}")
    print(f"Seasonality charts -> {SEASON_DIR}")
    print(f"Excel report -> {XLSX_PATH}")
    for col, label in ZONES.items():
        s = trends[col]["slope"]
        print(f"{label}: slope={s:.4f} kW/hour ({s*24:.1f} kW/day, {s*24*7:.1f} kW/week)")
